from bintang import Bintang
from pathlib import Path
TEST_DIR = Path(__file__).resolve().parents[1] / 'tests'

def test_bintang_table_excel():
    bt = Bintang("test_bt")
    bt.create_table("person")
    bt['person'].insert({'id':1,'name':'John','surname':'Smith','address':'1 Station St'})
    bt['person'].insert(['id','name','surname','hobby','address'],[2,'Jane','Brown','Digging','8 Parade Rd'])
    bt['person'].insert(['id','name','surname','Address'], [3,'Okie','Dokie', '7 Ocean Rd'])
    bt['person'].insert(('id','name','hobby','Address'),(4,'Maria','Digging','7 Heaven Ave'))
    bt['person'].insert(('id','name','hobby','Address'),(5,'Bing','Digging',None))

    bt.create_table("fishing_club")
    bt['fishing_club'].insert({'FirstName':'Ajes','LastName':'Freeman','Membership':'Active'})
    bt['fishing_club'].insert({'FirstName':'John','LastName':'Smith','Membership':'Active'})
    bt['fishing_club'].insert({'FirstName':'Jane','LastName':'Brown','Membership':'Active'})
    bt['fishing_club'].insert({'FirstName':'Nutmeg','LastName':'Spaniel','Membership':'Active'})
    bt['fishing_club'].insert({'FirstName':'Zekey','LastName':'Pokey','Membership':'Active'})

    # Create Excel workbooks for testing
    # export person table to xlsx
    xlsx_path = TEST_DIR / 'person.xlsx'
    xlsx_path.unlink(missing_ok=True) # delete existing if any

    from openpyxl import Workbook
    wb = Workbook()
    bt['person'].to_excel(wb, xlsx_path)
    assert xlsx_path.exists()

    # export fishing_club table to xls
    xls_path = TEST_DIR / 'fishing_club.xls'
    xls_path.unlink(missing_ok=True) # delete existing if any
    from xlwt import Workbook
    wbwt = Workbook()
    bt['fishing_club'].to_excel(wbwt, xls_path)
    assert xls_path.exists()


    # test reading back the Excel files
    from openpyxl import load_workbook
    wb_loaded = load_workbook(xlsx_path)
    ws = wb_loaded.active
    assert ws['A1'].value == 'id'
    assert ws['B1'].value == 'name'
    assert ws['C1'].value == 'surname'

    import xlrd
    wb_loaded_xls = xlrd.open_workbook(xls_path)
    ws_xls = wb_loaded_xls.sheet_by_index(0)
    assert ws_xls.cell_value(0,0) == 'FirstName'
    assert ws_xls.cell_value(0,1) == 'LastName'
    assert ws_xls.cell_value(0,2) == 'Membership'

    # test reading back the xlsx file using bintang table
    bt.create_table("person_from_excel")
    wb = load_workbook(xlsx_path)
    bt['person_from_excel'].read_excel(wb, 'person')
    assert len(bt['person_from_excel']) == len(bt['person'])

    # test reading back the xls file using bintang table
    bt.create_table("fishing_club_from_excel")
    import xlrd
    wb_xls = xlrd.open_workbook(xls_path)
    bt['fishing_club_from_excel'].read_excel(wb_xls, 'fishing_club')
    assert len(bt['fishing_club_from_excel']) == len(bt['fishing_club'])